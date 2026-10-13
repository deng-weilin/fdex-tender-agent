import copy
import json
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from fdex_tender.core import InputError, digest_file, digest_bytes, load_context, read_json, safe_path, validate_envelope
from fdex_tender.pipeline import run, release, rerun_module
from fdex_tender.modules import register, read_tender, response_matrix, material_index, match_materials, selection_fingerprint
from fdex_tender.documents import review_document

ROOT = Path(__file__).resolve().parents[1]


def load_m14_adapter():
    """按文件路径加载验收适配器 adapters/m14.py（它不是包，不能 import）。"""
    import importlib.util
    spec = importlib.util.spec_from_file_location("m14_adapter_under_test", ROOT / "adapters" / "m14.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="tender-test-")
        self.base = Path(self.temp.name)
        self.source = self.base / "input"
        shutil.copytree(ROOT / "samples/e2e/core-normal", self.source)

    def tearDown(self):
        self.temp.cleanup()

    def generate(self, scenario=None, name="run"):
        source = self.source if scenario is None else ROOT / "samples/e2e" / scenario
        target = self.base / name
        return target, run(source, target)

    def test_normal_actual_chain_and_schema(self):
        output, values = self.generate()
        self.assertTrue(all(v["processing_status"] == "COMPLETE" for v in values.values()))
        self.assertEqual(len(values["M04"]["data"]["requirements"]), 12)
        self.assertEqual(values["M07"]["data"]["coverage"]["scope_total"], 8)
        self.assertFalse(values["M07"]["data"]["coverage"]["full_tender_complete"])
        self.assertEqual(len(values["M09"]["data"]["selected"]), 3)
        for value in values.values():
            validate_envelope(value)
        self.assertTrue((output / "personnel-chapter.docx").exists())
        with zipfile.ZipFile(output / "archive.zip") as z:
            self.assertNotIn("outputs/.context.json", z.namelist())
            self.assertIn("inputs/context.json", z.namelist())

    def test_unconfirmed_does_not_generate_selection_docx(self):
        output, values = self.generate("unconfirmed")
        self.assertEqual(values["M09"]["processing_status"], "REVIEW")
        self.assertEqual(values["M09"]["data"]["selected"], [])
        self.assertFalse((output / "personnel-chapter.docx").exists())
        self.assertEqual(values["M13"]["data"]["artifact_status"], "NOT_GENERATED")

    def test_expired_and_changed_evidence_invalidates_confirmation(self):
        output, values = self.generate("expired")
        self.assertEqual(values["M09"]["data"]["confirmation"]["status"], "STALE")
        self.assertEqual(values["M09"]["data"]["candidates"][2]["status"], "EXPIRED")
        self.assertFalse((output / "personnel-chapter.docx").exists())

    def test_missing_image_not_fabricated(self):
        output, values = self.generate("missing-image")
        self.assertFalse((output / "personnel-chapter.docx").exists())
        self.assertTrue(any(i["code"] == "IMAGE_UNREADABLE" for i in values["M08"]["issues"]))

    def test_unselected_missing_image_does_not_block_selected_scope(self):
        (self.source / "materials/cert-04.png").unlink()
        output, values = self.generate()
        self.assertEqual(values["M14"]["processing_status"], "COMPLETE")
        self.assertTrue((output / "personnel-chapter.docx").exists())

    def test_changed_input_confirmation_cannot_reuse_old_release(self):
        output, _ = self.generate()
        record = read_json(self.source / "selection.json")
        record["confirmed_by"] = "Different DEMO actor"
        (self.source / "selection.json").write_text(json.dumps(record), encoding="utf-8")
        with self.assertRaises(InputError):
            release(output, "DEMO reviewer")

    def test_replacement_removes_old_name_number_and_image(self):
        old, normal = self.generate(name="normal")
        new, changed = self.generate("replacement", "replacement")
        with zipfile.ZipFile(new / "personnel-chapter.docx") as z:
            xml = z.read("word/document.xml").decode()
            hashes = {digest_bytes(z.read(n)) for n in z.namelist() if n.startswith("word/media/")}
        self.assertNotIn("示例人员丙", xml)
        self.assertNotIn("DEMO-CERT-003", xml)
        self.assertIn("示例人员丁", xml)
        self.assertNotIn(digest_file(self.source / "materials/cert-03.png"), hashes)
        self.assertIn(digest_file(self.source / "materials/cert-04.png"), hashes)

    def test_stable_ids_across_runs_and_changed_text(self):
        _, first = self.generate(name="first")
        text = self.source / "tender/tender.md"
        text.write_text(text.read_text().replace("人员不少于3名", "人员不少于4名"), encoding="utf-8")
        _, second = self.generate(name="second")
        a = [r["requirement_id"] for r in first["M04"]["data"]["requirements"]]
        b = [r["requirement_id"] for r in second["M04"]["data"]["requirements"]]
        self.assertEqual(a, b)
        self.assertEqual(second["M04"]["data"]["personnel_rule"]["count"], 4)
        self.assertEqual(second["M09"]["data"]["confirmation"]["status"], "STALE")

    def test_unknown_selected_id_is_rejected(self):
        record = read_json(self.source / "selection.json")
        record["person_ids"] = ["PERSON-01", "PERSON-02", "NOT-A-PERSON"]
        ctx = load_context(self.source, "TEST")
        r = register(ctx); t = read_tender(ctx, r); idx = material_index(ctx, r)
        record["target_digest"] = selection_fingerprint(ctx, t, idx, record["person_ids"], record["roles"])
        (self.source / "selection.json").write_text(json.dumps(record), encoding="utf-8")
        _, values = self.generate()
        self.assertEqual(values["M09"]["data"]["selected"], [])
        self.assertTrue(any(i["code"] == "SELECTION_INVALID" for i in values["M09"]["issues"]))

    def test_no_real_mode_fallback(self):
        context = read_json(self.source / "context.json")
        context["source_mode"] = "REAL"
        (self.source / "context.json").write_text(json.dumps(context), encoding="utf-8")
        with self.assertRaises(InputError):
            self.generate()
        self.assertFalse((self.base / "run").exists())

    def test_module_types_are_validated_not_just_field_presence(self):
        _, values = self.generate()
        broken = copy.deepcopy(values["M04"])
        broken["data"]["requirements"] = "not an array"
        with self.assertRaises(InputError):
            validate_envelope(broken)

    def test_nonpublic_source_never_creates_public_zip(self):
        path = self.source / "sources.json"
        data = read_json(path); data["sources"][0]["public"] = False
        path.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaises(InputError):
            self.generate()
        self.assertFalse((self.base / "run/archive.zip").exists())
        self.assertFalse(read_json(self.base / "run/archive_manifest.json")["data"]["public_export"])

    def test_path_traversal_and_symlink_escape(self):
        with self.assertRaises(InputError):
            safe_path(self.source, "../outside")
        (self.source / "outside-link").symlink_to(self.base)
        with self.assertRaises(InputError):
            safe_path(self.source, "outside-link/anything")

    def test_template_hash_or_columns_are_not_silently_accepted(self):
        path = self.source / "template-manifest.json"
        manifest = read_json(path); manifest["sha256"] = "0" * 64
        path.write_text(json.dumps(manifest), encoding="utf-8")
        output, values = self.generate()
        self.assertTrue(any(i["code"] == "TEMPLATE_HASH_MISMATCH" for i in values["M07"]["issues"]))
        self.assertFalse((output / "personnel-chapter.docx").exists())

    def test_cross_lot_confirm_rejected(self):
        record = read_json(self.source / "selection.json"); record["lot_id"] = "LOT-OTHER"
        (self.source / "selection.json").write_text(json.dumps(record), encoding="utf-8")
        _, values = self.generate()
        self.assertTrue(any(i["code"] == "CONFIRMATION_CONTEXT_MISMATCH" for i in values["M09"]["issues"]))

    def test_no_overwrite_of_history(self):
        self.generate()
        with self.assertRaises(InputError):
            self.generate()

    def test_review_detects_four_actual_document_faults(self):
        output, values = self.generate()
        path = output / "personnel-chapter.docx"
        with zipfile.ZipFile(path) as z:
            contents = {n: z.read(n) for n in z.namelist()}
        xml = contents["word/document.xml"].decode().replace("示例专业服务项目", "错误项目名").replace("示例人员丙", "旧人员姓名")
        xml = xml.replace("</w:body>", "<w:p><w:r><w:t>{{UNFILLED}}</w:t></w:r></w:p></w:body>")
        contents["word/document.xml"] = xml.encode()
        # Remove one visible image relationship target while preserving a valid DOCX package.
        media = next(n for n in contents if n.startswith("word/media/"))
        from lxml import etree
        rels = etree.fromstring(contents["word/_rels/document.xml.rels"])
        for node in list(rels):
            if node.get("Target") == media.removeprefix("word/"):
                rid = node.get("Id")
                rels.remove(node)
                xml = contents["word/document.xml"].decode()
                xml = re_sub_drawing(xml, rid)
                contents["word/document.xml"] = xml.encode()
        contents["word/_rels/document.xml.rels"] = etree.tostring(rels)
        contents.pop(media)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name, body in contents.items():
                z.writestr(name, body)
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], values["M13"], output)
        codes = {i["code"] for i in review["issues"]}
        self.assertTrue({"PROJECT_FIELDS_WRONG", "PERSON_TABLE_WRONG", "ATTACHMENT_SET_WRONG", "PLACEHOLDER_REMAINS"} <= codes)
        with self.assertRaises(InputError):
            release(output, "DEMO reviewer")

    def test_release_is_local_scope_and_input_changes_are_refused(self):
        output, _ = self.generate()
        result = release(output, "DEMO reviewer")
        self.assertEqual(result["scope"], "personnel_chapter_only")
        self.assertFalse(result["real_submission_authorized"])
        path = self.source / "materials/person-01.md"
        path.write_text(path.read_text() + "\nChanged", encoding="utf-8")
        with self.assertRaises(InputError):
            release(output, "DEMO reviewer")

    def test_module_rerun_uses_generated_upstream_and_revokes_old_release(self):
        old, _ = self.generate()
        release(old, "DEMO reviewer")
        new = self.base / "rerun"
        values = rerun_module("M13", old, new)
        self.assertEqual(values["M13"]["processing_status"], "COMPLETE")
        self.assertEqual(read_json(new / "release_record.json")["status"], "REVOKED")
        self.assertFalse((new / "archive.zip").exists())
        with self.assertRaises(InputError):
            release(new, "DEMO reviewer")

    def test_archive_can_replay_from_original_inputs(self):
        output, _ = self.generate()
        unpacked = self.base / "unpacked"
        with zipfile.ZipFile(output / "archive.zip") as z:
            z.extractall(unpacked)
        values = run(unpacked / "inputs", self.base / "replay")
        self.assertEqual(values["M14"]["processing_status"], "COMPLETE")

    def test_trace_records_actual_cell_and_relationship_locations(self):
        output, values = self.generate()
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(output / "personnel-chapter.docx")
        selected = {p["person_id"]: p for p in values["M09"]["data"]["selected"]}
        for entry in values["M13"]["data"]["trace"]:
            person = selected[entry["person_id"]]
            row = doc.tables[entry["table_index"]].rows[entry["row_index"]]
            self.assertIn(person["name"], row.cells[entry["name_column"]].text)
            self.assertIn(person["certificate_id"], row.cells[entry["certificate_column"]].text)
            blips = list(doc.paragraphs[entry["paragraph_index"]]._element.iter(qn("a:blip")))
            self.assertEqual(blips[0].get(qn("r:embed")), entry["relationship_id"])

    def test_swapped_images_are_located_even_when_media_set_is_unchanged(self):
        output, values = self.generate()
        from docx import Document
        from docx.oxml.ns import qn
        path = output / "personnel-chapter.docx"
        doc = Document(path)
        blips = list(doc.element.body.iter(qn("a:blip")))
        first, second = blips[0].get(qn("r:embed")), blips[1].get(qn("r:embed"))
        blips[0].set(qn("r:embed"), second)
        blips[1].set(qn("r:embed"), first)
        doc.save(path)
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], values["M13"], output)
        codes = {i["code"] for i in review["issues"]}
        self.assertIn("ATTACHMENT_PERSON_MISMATCH", codes)
        self.assertIn("image_relationship", review["data"]["located_issue_kinds"])
        self.assertTrue(any(m["kind"] == "image_relationship" for m in review["data"]["location_mismatches"]))

    def test_trace_pointing_to_wrong_cell_is_located(self):
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        assembled["data"]["trace"][0]["name_column"] = 2
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        self.assertIn("TRACE_CELL_MISMATCH", {i["code"] for i in review["issues"]})
        mismatch = next(m for m in review["data"]["location_mismatches"] if m["kind"] == "table_cell")
        self.assertEqual((mismatch["table_index"], mismatch["row_index"], mismatch["column_index"]), (0, 1, 2))

    def test_missing_trace_cannot_claim_location_review(self):
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        assembled["data"]["trace"] = []
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        self.assertEqual(review["processing_status"], "REVIEW")
        self.assertIn("TRACE_LOCATION_MISSING", {i["code"] for i in review["issues"]})

    def test_changed_document_invalidates_released_hash(self):
        output, values = self.generate()
        record = release(output, "DEMO reviewer")
        from docx import Document
        doc = Document(output / "personnel-chapter.docx")
        doc.paragraphs[0].add_run("（本地改写）")
        doc.save(output / "personnel-chapter.docx")
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], values["M13"], output, release=record)
        self.assertFalse(review["data"]["release_hash_matches"])
        self.assertIn("RELEASE_STALE", {i["code"] for i in review["issues"]})

    def test_placeholder_reports_the_exact_paragraph(self):
        output, values = self.generate()
        from docx import Document
        path = output / "personnel-chapter.docx"
        doc = Document(path)
        doc.paragraphs[0].text = "{{UNFILLED}}"
        doc.save(path)
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], values["M13"], output)
        self.assertIn("PLACEHOLDER_REMAINS", {i["code"] for i in review["issues"]})
        located = next(m for m in review["data"]["location_mismatches"]
                       if m["kind"] == "paragraph" and m.get("quote") == "{{UNFILLED}}")
        self.assertEqual(located["paragraph_index"], 0)
        self.assertEqual(located["found_in"], "paragraph")

    def test_person_table_mismatch_reports_the_differing_cell(self):
        output, values = self.generate()
        from docx import Document
        path = output / "personnel-chapter.docx"
        doc = Document(path)
        doc.tables[0].rows[1].cells[1].text = "未选用旧人员"
        doc.save(path)
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], values["M13"], output)
        self.assertIn("PERSON_TABLE_WRONG", {i["code"] for i in review["issues"]})
        located = next(m for m in review["data"]["location_mismatches"]
                       if m["kind"] == "table_cell" and m.get("field") is None and m.get("expected"))
        self.assertEqual((located["table_index"], located["row_index"], located["column_index"]), (0, 1, 1))
        self.assertEqual(located["actual"], "未选用旧人员")

    def test_missing_project_field_declares_the_scan_scope(self):
        output, values = self.generate()
        from docx import Document
        ctx = load_context(self.source, "TEST")
        path = output / "personnel-chapter.docx"
        doc = Document(path)
        body = [p for p in doc.paragraphs if ctx["project_name"] in p.text]
        self.assertTrue(body, "样例应有一处写入项目名")
        for paragraph in body:
            paragraph.text = ""
        doc.save(path)
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], values["M13"], output)
        self.assertIn("PROJECT_FIELDS_WRONG", {i["code"] for i in review["issues"]})
        located = next(m for m in review["data"]["location_mismatches"]
                       if m["kind"] == "paragraph" and m.get("field") == "project_name")
        self.assertIsNone(located["paragraph_index"])
        self.assertGreater(located["searched_paragraphs"], 0)

    # --- trace 声明完整性：缺字段/缺人/坐标非法不得放行，完整声明不得误杀 ---
    # 对应 Issue #14 v2 与维护者复审意见：覆盖门禁与位置检查必须由生产代码
    # 自己把守，教学 trace 缺字段不能成为跳过的理由。

    def test_complete_requirement_mapping_still_passes(self):
        output, values = self.generate()
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], values["M13"], output)
        self.assertEqual(review["data"]["requirement_coverage"], "DECLARED_AND_MATCHED")
        self.assertIn("scope_coverage", review["data"]["checks"])
        self.assertTrue(review["data"]["release_allowed"])

    def test_single_trace_entry_without_requirement_ids_cannot_pass(self):
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        del assembled["data"]["trace"][0]["requirement_ids"]
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        codes = {i["code"] for i in review["issues"]}
        self.assertIn("TRACE_REQUIREMENT_UNDECLARED", codes)
        self.assertEqual(review["data"]["requirement_coverage"], "NOT_DECLARED_IN_TRACE")
        self.assertIn(assembled["data"]["trace"][0]["person_id"],
                      review["data"]["requirement_undeclared_person_ids"])
        self.assertFalse(review["data"]["release_allowed"])
        self.assertNotIn("scope_coverage", review["data"]["checks"])

    def test_all_trace_entries_without_requirement_ids_cannot_pass(self):
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        for entry in assembled["data"]["trace"]:
            del entry["requirement_ids"]
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        self.assertIn("TRACE_REQUIREMENT_UNDECLARED", {i["code"] for i in review["issues"]})
        self.assertFalse(review["data"]["release_allowed"])
        self.assertNotIn("scope_coverage", review["data"]["checks"])

    def test_empty_requirement_list_cannot_hide_behind_set_comparison(self):
        # 每人的 requirement_ids 是同一份全量清单，只比集合时清空某一条，
        # 并集不变、wanted==actual 仍成立——必须逐条挡住空数组。
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        victim = assembled["data"]["trace"][0]["person_id"]
        assembled["data"]["trace"][0]["requirement_ids"] = []
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        self.assertIn("RESPONSE_TRACE_MISSING", {i["code"] for i in review["issues"]})
        self.assertIn(victim, review["data"]["requirement_empty_declared_person_ids"])
        self.assertFalse(review["data"]["release_allowed"])

    def test_dropping_one_persons_trace_cannot_claim_location_review(self):
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        dropped = assembled["data"]["trace"].pop()["person_id"]
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        codes = {i["code"] for i in review["issues"]}
        self.assertIn("TRACE_PERSON_MISSING", codes)
        self.assertFalse(review["data"]["trace_location_checked"])
        self.assertFalse(review["data"]["release_allowed"])
        self.assertNotIn("trace_cell_locations", review["data"]["checks"])
        self.assertTrue(any(m["person_id"] == dropped for m in review["data"]["trace_integrity"]))

    def test_duplicated_person_trace_is_reported(self):
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        assembled["data"]["trace"].append(copy.deepcopy(assembled["data"]["trace"][-1]))
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        self.assertIn("TRACE_PERSON_DUPLICATED", {i["code"] for i in review["issues"]})
        self.assertFalse(review["data"]["trace_location_checked"])
        self.assertFalse(review["data"]["release_allowed"])

    def test_missing_relationship_field_blocks_release(self):
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        del assembled["data"]["trace"][0]["relationship_id"]
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        self.assertIn("TRACE_RELATION_UNDECLARED", {i["code"] for i in review["issues"]})
        self.assertFalse(review["data"]["trace_location_checked"])
        self.assertFalse(review["data"]["release_allowed"])

    def test_missing_location_field_blocks_release(self):
        output, values = self.generate()
        assembled = copy.deepcopy(values["M13"])
        del assembled["data"]["trace"][0]["paragraph_index"]
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
        self.assertIn("TRACE_LOCATION_UNDECLARED", {i["code"] for i in review["issues"]})
        self.assertFalse(review["data"]["trace_location_checked"])
        self.assertFalse(review["data"]["release_allowed"])

    def test_negative_and_bool_coordinates_are_rejected(self):
        # bool 是 int 的子类，tables[True] 等价于 tables[1]；负数会静默取到
        # 末尾几行——两者都能让"指错位置"看起来像"指对了"。
        output, values = self.generate()
        ctx = load_context(self.source, "TEST")
        for field, bad in (("row_index", -1), ("table_index", True)):
            assembled = copy.deepcopy(values["M13"])
            assembled["data"]["trace"][0][field] = bad
            review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
            self.assertIn("TRACE_COORDINATE_INVALID", {i["code"] for i in review["issues"]},
                          "%s=%r 应被判坐标非法" % (field, bad))
            self.assertFalse(review["data"]["trace_location_checked"])
            self.assertFalse(review["data"]["release_allowed"])

    def test_complete_trace_still_claims_location_review(self):
        output, values = self.generate()
        ctx = load_context(self.source, "TEST")
        review = review_document(ctx, values["M07"], values["M08"], values["M09"], values["M13"], output)
        self.assertTrue(review["data"]["trace_location_checked"])
        self.assertIn("trace_cell_locations", review["data"]["checks"])
        self.assertIn("trace_image_relationships", review["data"]["checks"])
        self.assertTrue(review["data"]["release_allowed"])

    def test_adapter_mapping_fills_only_absent_keys(self):
        # 第三轮评审：补齐的唯一条件是 key 完全不存在。key 已存在时，无论值是
        # null、字符串、数字、空数组还是合法列表，一律原样保留交生产审查。
        adapter = load_m14_adapter()
        output, values = self.generate()
        in_scope = [r["requirement_id"] for r in values["M04"]["data"]["requirements"] if r.get("in_personnel_scope")]
        for label, declared in (("null", None), ("string", "REQ-FAKE"), ("number", 42),
                                ("empty_list", []), ("valid_list", list(in_scope))):
            trace = [{"person_id": "PERSON-01", "requirement_ids": declared},
                     {"person_id": "PERSON-02"}]
            filled = adapter.attach_requirement_mapping(trace, values["M04"])
            self.assertEqual(filled, ["PERSON-02"], "%s：只有缺 key 的条目才补齐" % label)
            self.assertEqual(trace[0]["requirement_ids"], declared, "%s：已声明的值绝不被覆盖" % label)
            self.assertEqual(trace[1]["requirement_ids"], in_scope, "%s：缺 key 的条目按生产规则补齐" % label)
        self.assertTrue(in_scope, "样例应有 in_personnel_scope 为真的要求")

    def test_declared_invalid_mapping_is_not_laundered_into_a_pass(self):
        # 第三轮评审的复现锁：null/字符串/数字曾被适配器洗成全量清单从而放行。
        # 经适配器后，这些值必须以原面目到达生产门禁并各自被对应 issue 打回。
        adapter = load_m14_adapter()
        output, values = self.generate()
        ctx = load_context(self.source, "TEST")
        for label, declared, expect_code in (("null", None, "TRACE_REQUIREMENT_UNDECLARED"),
                                             ("string", "REQ-FAKE", "TRACE_REQUIREMENT_UNDECLARED"),
                                             ("number", 42, "TRACE_REQUIREMENT_UNDECLARED"),
                                             ("empty_list", [], "RESPONSE_TRACE_MISSING")):
            assembled = copy.deepcopy(values["M13"])
            for entry in assembled["data"]["trace"]:
                entry.pop("requirement_ids", None)
            victim = assembled["data"]["trace"][0]["person_id"]
            assembled["data"]["trace"][0]["requirement_ids"] = declared
            filled = adapter.attach_requirement_mapping(assembled["data"]["trace"], values["M04"])
            self.assertNotIn(victim, filled, "%s：已声明的条目不得进补齐名单" % label)
            self.assertEqual(assembled["data"]["trace"][0]["requirement_ids"], declared,
                             "%s：值必须原样到达生产审查" % label)
            review = review_document(ctx, values["M07"], values["M08"], values["M09"], assembled, output)
            codes = {i["code"] for i in review["issues"]}
            self.assertIn(expect_code, codes, "%s 必须被生产门禁打回" % label)
            self.assertFalse(review["data"]["release_allowed"], "%s 不得放行" % label)

    def test_mapped_trace_passes_the_undeclared_gate_like_the_normal_case(self):
        # 正常专项用例的完整路径：fixture 裸 trace 先被门禁打回，适配器按上游
        # M04 映射补齐后同一份审查必须放行——锁定"适配器存在"这件事。
        adapter = load_m14_adapter()
        output, values = self.generate()
        ctx = load_context(self.source, "TEST")
        bare = copy.deepcopy(values["M13"])
        for entry in bare["data"]["trace"]:
            del entry["requirement_ids"]
        before = review_document(ctx, values["M07"], values["M08"], values["M09"], bare, output)
        self.assertIn("TRACE_REQUIREMENT_UNDECLARED", {i["code"] for i in before["issues"]})
        self.assertFalse(before["data"]["release_allowed"])

        filled = adapter.attach_requirement_mapping(bare["data"]["trace"], values["M04"])
        after = review_document(ctx, values["M07"], values["M08"], values["M09"], bare, output)
        self.assertEqual(filled, [e["person_id"] for e in bare["data"]["trace"]])
        self.assertEqual(after["processing_status"], "COMPLETE")
        self.assertEqual(after["data"]["requirement_coverage"], "DECLARED_AND_MATCHED")
        self.assertIn("scope_coverage", after["data"]["checks"])
        self.assertTrue(after["data"]["release_allowed"])
        self.assertEqual([i["code"] for i in after["issues"]], [])


def re_sub_drawing(xml, rid):
    from lxml import etree
    tree = etree.fromstring(xml.encode())
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    for drawing in tree.xpath("//w:drawing", namespaces=ns):
        if drawing.xpath(".//a:blip[@r:embed=$rid]", namespaces=ns, rid=rid):
            drawing.getparent().remove(drawing)
    return etree.tostring(tree, encoding="unicode")


if __name__ == "__main__":
    unittest.main()
