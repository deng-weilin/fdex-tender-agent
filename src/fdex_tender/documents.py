"""Editable personnel chapter and an independent review of actual OOXML."""
import re
import zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from .core import InputError, digest_file, digest_bytes, envelope, issue, require_same_context, safe_path


def _picture_relationship(shape):
    """取刚插入图片的关系ID，供 trace 记录它实际挂在哪个关系上。"""
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    return blip.get(qn("r:embed"))


def _person_image_blocks(document, anchors):
    """把可见图片按证书编号锚定归属到人。

    只看压缩包里的 media 集合，无法发现"图片都在文件里但挂错了人"；只看固定段落
    序号，又会因为文档里插入一段内容整体漂移而误报。因此归属规则是：图片前面最近
    一段含证书编号的正文，就是这张图片的责任人。
    """
    blocks = {person_id: [] for person_id in anchors.values()}
    current, unanchored = None, 0
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        for certificate_id, person_id in anchors.items():
            if certificate_id in paragraph.text:
                current = person_id
                break
        for blip in paragraph._element.iter(qn("a:blip")):
            relationship_id = blip.get(qn("r:embed"))
            if relationship_id is None or relationship_id not in document.part.rels:
                continue
            relation = document.part.rels[relationship_id]
            if relation.is_external or not relation.reltype.endswith("/image"):
                continue
            item = {"paragraph_index": paragraph_index, "relationship_id": relationship_id,
                    "image_sha256": digest_bytes(relation.target_part.blob)}
            if current is None:
                unanchored += 1
            else:
                blocks[current].append(item)
    return blocks, unanchored


def _verify_table_locations(document, index, selected, trace):
    """核对 trace 声明的表格单元格：姓名和证书编号必须真的写在那个格子里。"""
    assets = {a["asset_id"]: a for a in index["data"]["assets"]}
    by_id = {p["person_id"]: p for p in selected}
    findings = []
    for entry in trace:
        person = by_id.get(entry.get("person_id"))
        if person is None:
            findings.append((issue("TRACE_PERSON_UNKNOWN", str(entry.get("person_id")), "trace 登记了本次未确认的人员", owner="文档负责人"),
                             {"person_id": entry.get("person_id"), "kind": "table_cell", "detail": "trace 中的人员不在本次确认范围"}))
            continue
        asset = assets.get(person["asset_id"], {})
        for column_key, expected, label in (("name_column", person["name"], "姓名"),
                                            ("certificate_column", asset.get("certificate_id"), "证书编号")):
            column = entry.get(column_key)
            cell = None
            try:
                cell = document.tables[entry["table_index"]].rows[entry["row_index"]].cells[column]
            except (IndexError, KeyError, TypeError):
                cell = None
            if expected is None or cell is None or expected not in cell.text:
                findings.append((issue("TRACE_CELL_MISMATCH", f"{person['person_id']}-{label}",
                                       f"trace 声称{label}位于表格{entry.get('table_index')}第{entry.get('row_index')}行第{column}列，实际内容不符",
                                       owner="文档负责人"),
                                 {"person_id": person["person_id"], "kind": "table_cell",
                                  "table_index": entry.get("table_index"), "row_index": entry.get("row_index"),
                                  "column_index": column, "field": label,
                                  "expected": expected, "actual": None if cell is None else cell.text}))
    return findings


def _verify_image_relationships(document, index, selected, trace):
    """核对每个人的证书附件区：实际显示的图片必须属于这个人。"""
    assets = {a["asset_id"]: a for a in index["data"]["assets"]}
    anchors = {}
    for person in selected:
        asset = assets.get(person["asset_id"])
        if asset and asset.get("certificate_id"):
            anchors[asset["certificate_id"]] = person["person_id"]
    blocks, unanchored = _person_image_blocks(document, anchors)
    trace_by_person = {e.get("person_id"): e for e in trace}
    findings, drift = [], []
    for person in selected:
        asset = assets.get(person["asset_id"], {})
        expected = asset.get("sha256")
        shown = blocks.get(person["person_id"], [])
        entry = trace_by_person.get(person["person_id"], {})
        if not shown:
            findings.append((issue("ATTACHMENT_MISSING", person["person_id"], "实际Word中该人员的证书附件区没有可见图片", owner="文档负责人"),
                             {"person_id": person["person_id"], "kind": "image_relationship",
                              "detail": "附件区无可见图片", "expected_image_sha256": expected}))
            continue
        if len(shown) > 1:
            findings.append((issue("ATTACHMENT_PERSON_MISMATCH", person["person_id"], "同一人员的证书附件区出现多张图片", owner="文档负责人"),
                             {"person_id": person["person_id"], "kind": "image_relationship",
                              "detail": "附件区图片数量不是1", "shown": [i["relationship_id"] for i in shown]}))
        for item in shown:
            if item["image_sha256"] != expected:
                findings.append((issue("ATTACHMENT_PERSON_MISMATCH", person["person_id"], "证书附件区显示的图片与确认人员的证书不一致", owner="文档负责人"),
                                 {"person_id": person["person_id"], "kind": "image_relationship",
                                  "paragraph_index": item["paragraph_index"], "relationship_id": item["relationship_id"],
                                  "expected_image_sha256": expected, "actual_image_sha256": item["image_sha256"]}))
        if len(shown) == 1 and shown[0]["image_sha256"] == expected:
            claimed_rid = entry.get("relationship_id")
            if claimed_rid is not None and claimed_rid != shown[0]["relationship_id"]:
                findings.append((issue("ATTACHMENT_PERSON_MISMATCH", person["person_id"], "trace 声称的图片关系不是该人员附件区实际显示的关系", owner="文档负责人"),
                                 {"person_id": person["person_id"], "kind": "image_relationship",
                                  "detail": "trace 关系指向与实际显示不一致", "claimed_relationship_id": claimed_rid,
                                  "actual_relationship_id": shown[0]["relationship_id"]}))
        claimed = entry.get("image_sha256")
        if claimed is not None and claimed != expected:
            findings.append((issue("ATTACHMENT_PERSON_MISMATCH", person["person_id"], "trace 登记的图片哈希与确认证书不一致", owner="文档负责人"),
                             {"person_id": person["person_id"], "kind": "image_relationship",
                              "detail": "trace 图片哈希与确认证书不一致", "expected_image_sha256": expected, "claimed_image_sha256": claimed}))
        actual_index = shown[0]["paragraph_index"] if shown else None
        if entry.get("paragraph_index") is not None and actual_index is not None and entry["paragraph_index"] != actual_index:
            drift.append({"person_id": person["person_id"], "claimed_paragraph_index": entry["paragraph_index"],
                          "actual_paragraph_index": actual_index})
    return findings, drift, unanchored


def assemble(ctx, tender, matrix, index, matched, run_dir):
    require_same_context(ctx, tender, matrix, index, matched)
    problems = list(matrix["issues"]) + list(matched["issues"])
    result = {"artifact_status": "NOT_GENERATED", "files": [], "trace": [], "selected_person_ids": [], "selection_digest": matched["data"]["selection_digest"]}
    if not matched["data"]["selected"] or matched["data"]["confirmation"]["status"] != "CONFIRMED" or any(i["severity"] == "BLOCKER" for i in problems):
        return envelope(ctx, "M13", result, matrix["input_refs"], problems,
                        status="BLOCKED" if any(i["severity"] == "BLOCKER" for i in problems) else "REVIEW")
    manifest = matrix["data"]["template"]
    template = safe_path(ctx["input_root"], manifest["relative_path"])
    if digest_file(template) != manifest["sha256"]:
        problems.append(issue("TEMPLATE_HASH_MISMATCH", manifest["template_id"], "装配时模板哈希已变化"))
        return envelope(ctx, "M13", result, matrix["input_refs"], problems)
    doc = Document(template)
    replacements = {"{{PROJECT_NAME}}": ctx["project_name"], "{{PROJECT_ID}}": ctx["project_id"], "{{LOT_ID}}": ctx["lot_id"]}
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for key, value in replacements.items():
                run.text = run.text.replace(key, value)
    if len(doc.tables) != 1 or [cell.text for cell in doc.tables[0].rows[0].cells] != manifest["columns"]:
        return envelope(ctx, "M13", result, matrix["input_refs"], [issue("TEMPLATE_TABLE_MISMATCH", "template", "模板表格与字段映射不一致")])
    table = doc.tables[0]
    assets = {a["asset_id"]: a for a in index["data"]["assets"]}
    trace = []
    for ordinal, person in enumerate(matched["data"]["selected"], 1):
        asset = assets[person["asset_id"]]
        image_path = safe_path(ctx["input_root"], asset["relative_path"])
        if not image_path.is_file() or digest_file(image_path) != asset["sha256"]:
            return envelope(ctx, "M13", result, matrix["input_refs"], [issue("ASSET_HASH_CHANGED", asset["asset_id"], "确认后证书原件发生变化，需重新确认")])
        cells = table.add_row().cells
        values = [person["role"], person["name"], asset["certificate_id"], asset["valid_until"], asset["relative_path"]]
        for cell, value in zip(cells, values):
            cell.text = value
        doc.add_page_break()
        doc.add_heading(f"人员证书 {ordinal}", level=1)
        doc.add_paragraph(f"{person['name']}  {asset['certificate_id']}")
        doc.add_paragraph("教学示例证书，仅用于合成材料验证。")
        doc.add_picture(str(image_path), width=Inches(6.1))
        trace.append({"section_id": "SECTION-PERSONNEL", "field_id": f"PERSONNEL-ROW-{ordinal}",
                      "person_id": person["person_id"], "asset_id": person["asset_id"],
                      "attachment_id": f"ATT-{person['asset_id']}", "image_sha256": asset["sha256"],
                      "requirement_ids": [r["requirement_id"] for r in tender["data"]["requirements"] if r["in_personnel_scope"]],
                      "output_file": "personnel-chapter.docx", "output_location": f"人员表第{ordinal}行及证书附件{ordinal}",
                      "table_index": 0, "row_index": len(table.rows) - 1,
                      "name_column": 1, "certificate_column": 2,
                      "paragraph_index": len(doc.paragraphs) - 1,
                      "relationship_id": _picture_relationship(doc.inline_shapes[-1])})
    output = Path(run_dir) / "personnel-chapter.docx"
    doc.core_properties.author = "FDEX Teaching Baseline"
    doc.core_properties.title = "人员资质章节"
    doc.core_properties.subject = "DEMO personnel chapter only"
    doc.save(output)
    result.update(artifact_status="PARTIAL_DRAFT", files=[{"relative_path": output.name, "sha256": digest_file(output), "editable": True}],
                  trace=trace, selected_person_ids=[p["person_id"] for p in matched["data"]["selected"]],
                  template_id=manifest["template_id"], template_version=manifest["version"])
    return envelope(ctx, "M13", result, matrix["input_refs"] + index["input_refs"], problems)


def review_document(ctx, matrix, index, matched, assembled, run_dir, release=None):
    """独立审查实际Word，不只读装配日志。

    除项目字段、占位符、人员表和图片集合外，还把 trace 对应到实际表格单元格与
    附件关系：姓名/证号必须真的写在 trace 声称的那一格，每个人附件区实际显示的
    图片必须属于这个人。release 为已放行记录时，实际文件与放行哈希不一致即判旧
    放行失效。
    """
    require_same_context(ctx, matrix, index, matched, assembled)
    problems, checks = [], []
    located = {}
    coverage = matrix["data"]["coverage"]
    result = {"review_status": "DRAFT", "file_sha256": None, "checks": checks, "coverage": coverage}

    def record(problem, kind=None, mismatch=None):
        problems.append(problem)
        if kind:
            located.setdefault(kind, []).append(mismatch if mismatch is not None else {"kind": kind, "object_id": problem["object_id"]})

    if not assembled["data"]["files"]:
        return envelope(ctx, "M14", result, assembled["input_refs"], [issue("NO_DOCUMENT", "document", "没有可审查的实际Word，前序确认或材料问题待处理")])
    path = safe_path(run_dir, assembled["data"]["files"][0]["relative_path"])
    if not path.exists():
        return envelope(ctx, "M14", result, assembled["input_refs"], [issue("DOCUMENT_MISSING", "document", "实际Word不存在")])
    result["file_sha256"] = digest_file(path)
    if result["file_sha256"] != assembled["data"]["files"][0]["sha256"]:
        problems.append(issue("DOCUMENT_HASH_CHANGED", "document", "Word与装配记录版本不一致"))
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8")
            media = {digest_bytes(z.read(name)) for name in z.namelist() if name.startswith("word/media/")}
        doc = Document(path)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile):
        return envelope(ctx, "M14", result, assembled["input_refs"], [issue("DOCX_UNREADABLE", "document", "实际Word包损坏或关联图片丢失")])
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    for field, value in (("project_name", ctx["project_name"]), ("project_id", ctx["project_id"]), ("lot_id", ctx["lot_id"])):
        if value not in text:
            record(issue("PROJECT_FIELDS_WRONG", "document.title", f"实际文档未包含{field}对应内容"),
                   "paragraph", {"kind": "paragraph", "field": field, "expected": value, "detail": "实际文档缺少该项目字段"})
    placeholder = re.search(r"\{\{[^{}]+\}\}|\[TODO[^\]]*\]", text)
    if placeholder:
        record(issue("PLACEHOLDER_REMAINS", "document", "实际文档仍含未处理占位符"),
               "paragraph", {"kind": "paragraph", "quote": placeholder.group(0), "detail": "未处理占位符仍留在实际Word中"})
    assets = {a["asset_id"]: a for a in index["data"]["assets"]}
    selected = matched["data"]["selected"]
    table_rows = [[c.text for c in row.cells] for row in doc.tables[0].rows[1:]] if len(doc.tables) == 1 else []
    expected_rows = [[p["role"], p["name"], p["certificate_id"], p["valid_until"], assets[p["asset_id"]]["relative_path"]] for p in selected]
    if table_rows != expected_rows:
        record(issue("PERSON_TABLE_WRONG", "PERSONNEL-TABLE", "实际人员表与确认记录不一致，可能包含旧人员或错证书"),
               "table_cell", {"kind": "table_cell", "detail": "实际人员表与确认记录不一致",
                              "actual_rows": table_rows, "expected_rows": expected_rows})
    expected_media = {assets[p["asset_id"]]["sha256"] for p in selected}
    if media != expected_media:
        problems.append(issue("ATTACHMENT_SET_WRONG", "PERSONNEL-ATTACHMENTS", "实际图片缺失、错配或包含旧版本图片"))
    unselected = [p for p in index["data"]["persons"] if p["person_id"] not in {x["person_id"] for x in selected}]
    for person in unselected:
        old_assets = [a for a in index["data"]["assets"] if a["person_id"] == person["person_id"]]
        if person["name"] in text or any(a["certificate_id"] in text or a["sha256"] in media for a in old_assets):
            record(issue("UNSELECTED_PERSON_REMAINS", person["person_id"], "未选人员的姓名、编号或图片残留"),
                   "paragraph", {"kind": "paragraph", "person_id": person["person_id"], "detail": "未选人员信息残留在实际Word中"})
    trace = assembled["data"].get("trace") or []
    drift, unanchored = [], None
    if not trace:
        record(issue("TRACE_LOCATION_MISSING", "trace", "装配未提供位置记录，不能声称表格单元格与附件关系已定位检查", severity="MAJOR", owner="文档负责人"))
    else:
        for problem, mismatch in _verify_table_locations(doc, index, selected, trace):
            record(problem, "table_cell", mismatch)
        image_findings, drift, unanchored = _verify_image_relationships(doc, index, selected, trace)
        for problem, mismatch in image_findings:
            record(problem, "image_relationship", mismatch)
        if all("requirement_ids" in entry for entry in trace):
            rows = matrix["data"]["responses"]
            wanted = {rid for row in rows for rid in row["requirement_ids"]}
            actual = {rid for row in trace for rid in row.get("requirement_ids", [])}
            if wanted != actual:
                problems.append(issue("RESPONSE_TRACE_MISSING", "trace", "声明范围内要求到实际输出的映射缺失"))
        else:
            result["requirement_coverage"] = "NOT_DECLARED_IN_TRACE"
    inspected = result["file_sha256"]
    release_hash_matches = None
    if release:
        release_hash_matches = release.get("sha256") == inspected
        if release.get("status") == "RELEASED" and not release_hash_matches:
            record(issue("RELEASE_STALE", "release", "实际Word与已放行版本不一致，旧放行失效，需重新审查后再放行", owner="招投标负责人"))
    checks.extend(["actual_docx_opened", "project_fields", "selected_table", "embedded_image_hashes", "no_unselected_person",
                   "no_placeholder", "scope_coverage", "trace_cell_locations", "trace_image_relationships"])
    location_mismatches = [m for kind in ("table_cell", "image_relationship", "paragraph") for m in located.get(kind, [])]
    result.update({"location_mismatches": location_mismatches,
                   "located_issue_kinds": [kind for kind in ("table_cell", "image_relationship", "paragraph") if located.get(kind)],
                   "inspected_document_sha256": inspected,
                   "trace_location_checked": bool(trace),
                   "paragraph_index_drift": drift,
                   "unanchored_images": unanchored,
                   "release_allowed": not problems and matched["data"]["confirmation"]["status"] == "CONFIRMED"})
    if release_hash_matches is not None:
        result["release_hash_matches"] = release_hash_matches
    result["review_status"] = "CHANGES_REQUIRED" if problems else "DRAFT"
    result["visual_review"] = "NOT_PERFORMED_BY_PROGRAM"
    return envelope(ctx, "M14", result, assembled["input_refs"], problems)
